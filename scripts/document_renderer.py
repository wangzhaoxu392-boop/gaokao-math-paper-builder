#!/usr/bin/env python3
"""Render a selected paper into DOCX/PDF and generate deterministic QA artifacts."""

from __future__ import annotations

import html
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parent.parent
BAND_DIRS = ["L1_30", "L2_60", "L3_90", "L4_120", "L5_120plus"]
SERIF_CN = Path(r"C:\Windows\Fonts\simfang.ttf")
HEI_CN = Path(r"C:\Windows\Fonts\simhei.ttf")
NOTO_SERIF = Path(r"C:\Windows\Fonts\NotoSerifSC-VF.ttf")
NOTO_SANS = Path(r"C:\Windows\Fonts\NotoSansSC-VF.ttf")
TYPE_LABEL = {
    "single_choice": "单选题",
    "multiple_choice": "多选题",
    "fill_blank": "填空题",
    "solution": "解答题",
}


class RenderError(RuntimeError):
    pass


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def find_item(item_id: str) -> dict[str, Any]:
    for band_dir in BAND_DIRS:
        candidate = ROOT / "bank" / "items" / band_dir / f"{item_id}.json"
        if candidate.exists():
            return load_json(candidate)
    for path in (ROOT / "bank" / "items").rglob("*.json"):
        item = load_json(path)
        if item.get("id") == item_id:
            return item
    raise RenderError(f"找不到题目：{item_id}")


def normalized(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def latex_fallback(latex: str) -> str:
    value = latex.strip()
    replacements = {
        r"\times": "×",
        r"\cdot": "·",
        r"\le": "≤",
        r"\ge": "≥",
        r"\neq": "≠",
        r"\pm": "±",
        r"\infty": "∞",
        r"\pi": "π",
        r"\alpha": "α",
        r"\beta": "β",
        r"\theta": "θ",
        r"\triangle": "△",
        r"\angle": "∠",
        r"\parallel": "∥",
        r"\perp": "⊥",
        r"\in": "∈",
        r"\notin": "∉",
        r"\subset": "⊂",
        r"\cup": "∪",
        r"\cap": "∩",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    value = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", value)
    value = re.sub(r"\\sqrt\{([^{}]+)\}", r"√(\1)", value)
    superscripts = str.maketrans("0123456789+-n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ⁿ")
    value = re.sub(r"\^\{?([0-9+\-n]+)\}?", lambda m: m.group(1).translate(superscripts), value)
    value = re.sub(r"\\[A-Za-z]+", "", value)
    value = value.replace("{", "").replace("}", "").replace("\\", "")
    return value


def find_poppler_tool(name: str) -> str | None:
    executable = shutil.which(name)
    if executable:
        return executable
    base = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "native" / "poppler"
    matches = sorted(base.rglob(f"{name}.exe")) if base.exists() else []
    return str(matches[0]) if matches else None


def render_latex_png(latex: str, path: Path, display: bool = False) -> tuple[str, float] | None:
    try:
        from matplotlib import mathtext  # type: ignore

        path.parent.mkdir(parents=True, exist_ok=True)
        mathtext.math_to_image(
            f"${latex}$",
            str(path),
            dpi=220,
            format="png",
        )
        try:
            from PIL import Image  # type: ignore

            with Image.open(path) as image:
                width, height = image.size
            return str(path), width / max(height, 1)
        except Exception:
            return str(path), 2.5 if display else 1.8
    except Exception:
        return None


FORMULA_PATTERN = re.compile(
    r"(?<!\\)\$(.+?)(?<!\\)\$|!\[[^\]]*\]\(([^)]+)\)",
    re.DOTALL,
)


def split_segments(text: str) -> list[tuple[str, str]]:
    parts: list[tuple[str, str]] = []
    cursor = 0
    for match in FORMULA_PATTERN.finditer(text or ""):
        if match.start() > cursor:
            parts.append(("text", text[cursor:match.start()]))
        if match.group(1) is not None:
            parts.append(("math", match.group(1).strip()))
        else:
            parts.append(("image", match.group(2).strip()))
        cursor = match.end()
    if cursor < len(text or ""):
        parts.append(("text", text[cursor:]))
    return parts or [("text", text or "")]


def add_rich_docx_paragraph(document: Any, text: str, style: str | None = None) -> Any:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches

    paragraph = document.add_paragraph(style=style)
    for kind, value in split_segments(text):
        if kind == "text":
            if value:
                paragraph.add_run(value)
            continue
        if kind == "image":
            image_path = (ROOT / value).resolve()
            if image_path.exists() and image_path.is_file():
                try:
                    from PIL import Image  # type: ignore

                    with Image.open(image_path) as image:
                        width, height = image.size
                    display_width = min(5.7, max(0.08, 0.17 * width / max(height, 1)))
                    paragraph.add_run().add_picture(
                        str(image_path), width=Inches(display_width)
                    )
                except Exception:
                    paragraph.add_run("[图片]")
            continue
        temp_path = Path(tempfile.mkdtemp(prefix="gm-formula-")) / "formula.png"
        rendered = render_latex_png(value, temp_path)
        if rendered:
            _, ratio = rendered
            height = Inches(0.18)
            width = Inches(min(5.8, max(0.18, 0.18 * ratio)))
            paragraph.add_run().add_picture(str(temp_path), width=width, height=height)
        else:
            paragraph.add_run(latex_fallback(value))
    if not split_segments(text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def configure_docx(document: Any) -> None:
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    styles = document.styles
    styles["Normal"].font.name = "FangSong"
    styles["Normal"]._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "FangSong",
    )
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.35
    for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "SimHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = footer.add_run()
    run._r.extend([begin, instruction, separate, placeholder, end])
    footer.add_run(" 页")


def add_docx_header_details(document: Any, request: dict[str, Any], blueprint: dict[str, Any], student: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("学生试卷" if student else "教师答案与教学包")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        f"{blueprint.get('name', '')} | 满分 {blueprint.get('total_score')} 分 | "
        f"建议 {blueprint.get('paper_minutes')} 分钟 | 当前 {request.get('current_band')} | "
        f"目标 {request.get('target_band')}"
    )
    if student:
        info = document.add_paragraph("姓名：____________    日期：____________    得分：____________")
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER


def create_docx(
    output_path: Path,
    manifest: dict[str, Any],
    items: list[dict[str, Any]],
    student: bool,
) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_BREAK
        from docx.shared import Pt
    except ImportError as exc:
        raise RenderError("缺少 python-docx，请先运行 gui/setup.ps1") from exc

    document = Document()
    configure_docx(document)
    add_docx_header_details(document, manifest["request"], manifest["blueprint"], student)
    document.add_paragraph()
    for slot in manifest["selected"]:
        item = item_by_id(items, slot["item_id"])
        heading = document.add_paragraph()
        run = heading.add_run(f"{slot['index']}. （{TYPE_LABEL[slot['question_type']]}，{slot['score']} 分）")
        run.bold = True
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.keep_together = True
        stem = add_rich_docx_paragraph(document, item.get("stem_markdown", ""))
        stem.paragraph_format.keep_with_next = True
        stem.paragraph_format.keep_together = True
        options = item.get("options", [])
        for option_index, option in enumerate(options):
            option_paragraph = add_rich_docx_paragraph(document, f"{chr(65 + option_index)}. {option}")
            option_paragraph.paragraph_format.keep_with_next = True
            option_paragraph.paragraph_format.keep_together = True
        for figure in item.get("figures", []):
            figure_path = (ROOT / figure).resolve()
            if figure_path.exists() and figure_path.is_file():
                document.add_picture(str(figure_path))
        if student:
            lines = 3 if slot["question_type"] != "solution" else 9
            for _ in range(lines):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.add_run(" ")
        else:
            answer_heading = document.add_paragraph()
            answer_heading.add_run("答案与解析").bold = True
            add_rich_docx_paragraph(document, item.get("answer", ""))
            if item.get("solution"):
                add_rich_docx_paragraph(document, item.get("solution", ""))
            rubric = item.get("rubric", [])
            if rubric:
                rubric_heading = document.add_paragraph()
                rubric_heading.add_run("评分点").bold = True
                for rubric_point in rubric:
                    add_rich_docx_paragraph(document, f"- {rubric_point}")
        document.add_paragraph()
    if not student:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph(style="Heading 1").add_run("教学实施建议")
        request = manifest["request"]
        if request.get("delivery_type") == "tutoring":
            guidance = [
                "0-10 分钟：复查近期错误并明确本次目标。",
                "10-30 分钟：梳理公式、条件与标准方法。",
                "30-85 分钟：学生独立完成核心卷。",
                "85-110 分钟：按评分点讲评并订正。",
                "110-120 分钟：归纳错因并布置同考点变式任务。",
            ]
        elif request.get("delivery_type") == "classroom":
            guidance = [
                "0-5 分钟：旧知回顾与目标说明。",
                "5-10 分钟：方法提示和易错点提醒。",
                "10-40 分钟：学生独立完成小卷。",
                "40-45 分钟：自检、收卷和后置讲评安排。",
            ]
        else:
            guidance = ["按标准考试要求组织，完成后依据评分点讲评。"]
        for line in guidance:
            add_rich_docx_paragraph(document, line)
    document.save(output_path)


def xml_escape_with_math(text: str, cache_dir: Path, display: bool = False) -> str:
    values: list[str] = []
    for kind, value in split_segments(text):
        if kind == "text":
            values.append(html.escape(value).replace("\n", "<br/>"))
            continue
        if kind == "image":
            image_path = (ROOT / value).resolve()
            if image_path.exists() and image_path.is_file():
                try:
                    from PIL import Image  # type: ignore

                    with Image.open(image_path) as image:
                        width, height = image.size
                    display_height = 16
                    display_width = max(6, min(260, display_height * width / max(height, 1)))
                    values.append(
                        f'<img src="{html.escape(str(image_path))}" width="{display_width:.0f}" height="{display_height}" valign="-4"/>'
                    )
                except Exception:
                    values.append("[图片]")
            continue
        safe = re.sub(r"[^A-Za-z0-9_-]", "_", value)[:50]
        formula_path = cache_dir / f"f_{len(list(cache_dir.glob('f_*.png'))):04d}_{safe}.png"
        rendered = render_latex_png(value, formula_path, display)
        if rendered:
            path, ratio = rendered
            height = 16 if not display else 22
            width = max(8, min(260, height * ratio))
            values.append(
                f'<img src="{html.escape(path)}" width="{width:.0f}" height="{height}" valign="-4"/>'
            )
        else:
            values.append(html.escape(latex_fallback(value)))
    return "".join(values)


def create_pdf(
    output_path: Path,
    manifest: dict[str, Any],
    items: list[dict[str, Any]],
    student: bool,
) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            KeepTogether,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )
    except ImportError as exc:
        raise RenderError("缺少 reportlab，请先运行 gui/setup.ps1") from exc

    body_font = SERIF_CN if SERIF_CN.exists() else NOTO_SERIF
    bold_font = HEI_CN if HEI_CN.exists() else NOTO_SANS
    pdfmetrics.registerFont(TTFont("CNBody", str(body_font)))
    pdfmetrics.registerFont(TTFont("CNBold", str(bold_font)))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CNTitle",
        parent=styles["Title"],
        fontName="CNBold",
        fontSize=20,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.black,
        spaceAfter=8,
    )
    meta_style = ParagraphStyle(
        "CNMeta",
        parent=styles["Normal"],
        fontName="CNBody",
        fontSize=10,
        leading=17,
        alignment=TA_CENTER,
        textColor=colors.black,
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "CNBody",
        parent=styles["Normal"],
        fontName="CNBody",
        fontSize=10.5,
        leading=17,
        textColor=colors.black,
        spaceAfter=5,
    )
    heading_style = ParagraphStyle(
        "CNHeading",
        parent=body_style,
        fontName="CNBold",
        fontSize=12,
        leading=18,
        spaceBefore=8,
        spaceAfter=5,
    )
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="学生试卷" if student else "教师答案与教学包",
        author="",
    )
    story: list[Any] = []
    cache_dir = output_path.parent / ".formula-cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    request = manifest["request"]
    blueprint = manifest["blueprint"]
    story.append(Paragraph("学生试卷" if student else "教师答案与教学包", title_style))
    story.append(
        Paragraph(
            html.escape(
                f"{blueprint.get('name', '')} | 满分 {blueprint.get('total_score')} 分 | "
                f"建议 {blueprint.get('paper_minutes')} 分钟 | 当前 {request.get('current_band')} | "
                f"目标 {request.get('target_band')}"
            ),
            meta_style,
        )
    )
    if student:
        story.append(Paragraph("姓名：____________    日期：____________    得分：____________", meta_style))
    for slot in manifest["selected"]:
        item = item_by_id(items, slot["item_id"])
        block: list[Any] = []
        block.append(
            Paragraph(
                html.escape(
                    f"{slot['index']}. （{TYPE_LABEL[slot['question_type']]}，{slot['score']} 分）"
                ),
                heading_style,
            )
        )
        block.append(Paragraph(xml_escape_with_math(item.get("stem_markdown", ""), cache_dir), body_style))
        for option_index, option in enumerate(item.get("options", [])):
            block.append(
                Paragraph(
                    f"{chr(65 + option_index)}. {xml_escape_with_math(str(option), cache_dir)}",
                    body_style,
                )
            )
        if student:
            block.append(Spacer(1, 8 * mm if slot["question_type"] != "solution" else 25 * mm))
        else:
            block.append(Paragraph("答案与解析", heading_style))
            block.append(Paragraph(xml_escape_with_math(item.get("answer", ""), cache_dir), body_style))
            if item.get("solution"):
                block.append(Paragraph(xml_escape_with_math(item.get("solution", ""), cache_dir), body_style))
            if item.get("rubric"):
                block.append(Paragraph("评分点", heading_style))
                for point in item["rubric"]:
                    block.append(Paragraph("• " + xml_escape_with_math(point, cache_dir), body_style))
        story.append(KeepTogether(block))
    if not student:
        story.append(PageBreak())
        story.append(Paragraph("教学实施建议", heading_style))
        if request.get("delivery_type") == "tutoring":
            timeline = [
                "0-10 分钟：复查近期错误并明确本次目标。",
                "10-30 分钟：梳理公式、条件与标准方法。",
                "30-85 分钟：学生独立完成核心卷。",
                "85-110 分钟：按评分点讲评并订正。",
                "110-120 分钟：归纳错因并布置同考点变式任务。",
            ]
        elif request.get("delivery_type") == "classroom":
            timeline = [
                "0-5 分钟：旧知回顾与目标说明。",
                "5-10 分钟：方法提示和易错点提醒。",
                "10-40 分钟：学生独立完成小卷。",
                "40-45 分钟：自检、收卷和后置讲评安排。",
            ]
        else:
            timeline = ["按标准考试要求组织，完成后依据评分点讲评。"]
        for line in timeline:
            story.append(Paragraph(html.escape(line), body_style))
    def footer(canvas: Any, current_doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("CNBody", 9)
        canvas.drawCentredString(A4[0] / 2, 10 * mm, f"第 {current_doc.page} 页")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def item_by_id(items: list[dict[str, Any]], item_id: str) -> dict[str, Any]:
    for item in items:
        if item.get("id") == item_id:
            return item
    raise RenderError(f"选题清单引用了不存在的题目：{item_id}")


def student_markdown(manifest: dict[str, Any], items: list[dict[str, Any]]) -> str:
    request = manifest["request"]
    blueprint = manifest["blueprint"]
    lines = [
        f"# {blueprint.get('name', '数学试卷')}",
        "",
        f"- 交付类型：{request.get('delivery_type')}",
        f"- 当前档位：{request.get('current_band')}",
        f"- 目标档位：{request.get('target_band')}",
        f"- 薄弱专题：{', '.join(request.get('weak_topics', []))}",
        f"- 总分：{blueprint.get('total_score')}",
        f"- 建议时长：{blueprint.get('paper_minutes')} 分钟",
        f"- 随机种子：{manifest.get('seed')}",
        "",
        "## 学生卷",
        "",
    ]
    for slot in manifest["selected"]:
        item = item_by_id(items, slot["item_id"])
        lines.extend(
            [
                f"### {slot['index']}. {TYPE_LABEL[slot['question_type']]}（{slot['score']}分）",
                "",
                item.get("stem_markdown", ""),
                "",
            ]
        )
        for index, option in enumerate(item.get("options", [])):
            lines.append(f"{chr(65 + index)}. {option}")
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def teacher_markdown(manifest: dict[str, Any], items: list[dict[str, Any]]) -> str:
    lines = ["# 教师答案与教学建议", ""]
    for slot in manifest["selected"]:
        item = item_by_id(items, slot["item_id"])
        lines.extend(
            [
                f"### {slot['index']}. 答案与解析",
                "",
                item.get("answer", ""),
                "",
                item.get("solution", ""),
                "",
            ]
        )
        if item.get("rubric"):
            lines.append("评分点：")
            lines.extend(f"- {point}" for point in item["rubric"])
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def render_pdf_pngs(pdf_path: Path, qa_dir: Path) -> list[str]:
    executable = find_poppler_tool("pdftoppm")
    if not executable:
        return []
    qa_dir.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [executable, "-png", "-r", "120", str(pdf_path), str(qa_dir / "page")],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if completed.returncode != 0:
        raise RenderError(f"PDF 渲染失败：{completed.stderr[-500:]}")
    return [str(path) for path in sorted(qa_dir.glob("page-*.png"))]


def validate_rendered(
    manifest: dict[str, Any],
    items: list[dict[str, Any]],
    student_markdown: str,
) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    selected_total = sum(float(slot.get("score", 0)) for slot in manifest.get("selected", []))
    if abs(selected_total - float(manifest.get("total_score", selected_total))) > 0.001:
        errors.append("选题分值之和与总分不一致")
    for slot in manifest.get("selected", []):
        item = item_by_id(items, slot["item_id"])
        if item.get("answer_status") not in {"matched", "verified"}:
            errors.append(f"{item.get('id')} 答案状态不可用于正式卷")
        if item.get("review_status") != "verified":
            errors.append(f"{item.get('id')} 尚未通过复核")
        answer = normalized(item.get("answer", ""))
        if len(answer) >= 8 and answer in student_markdown:
            errors.append(f"{item.get('id')} 的答案疑似泄漏到学生卷")
        solution = normalized(item.get("solution", ""))
        if len(solution) >= 20 and solution in student_markdown:
            errors.append(f"{item.get('id')} 的解析疑似泄漏到学生卷")
    estimated = float(manifest.get("estimated_minutes", 0))
    allowed = float(manifest["blueprint"].get("paper_minutes", 0))
    if estimated > allowed * 1.10:
        errors.append(f"预计用时 {estimated:.1f} 分钟超过限制 {allowed:.0f} 分钟")
    if len(manifest.get("selected", [])) != int(manifest.get("question_count", 0)):
        errors.append("选题数量与清单不一致")
    if not manifest.get("item_ids"):
        errors.append("选题清单为空")
    if not any(ROOT.glob("bank/items/**/*.json")):
        warnings.append("题库为空")
    return {"ok": not errors, "errors": errors, "warnings": warnings}


def render_selection(selection_path: Path, output_dir: Path) -> dict[str, Any]:
    selection_path = selection_path.resolve()
    manifest = load_json(selection_path)
    output_dir = output_dir.resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise RenderError(f"输出目录已存在且不为空：{output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)
    items = [find_item(item_id) for item_id in manifest.get("item_ids", [])]
    manifest = {
        **manifest,
        "rendered_at": datetime_now(),
        "selection_file": str(selection_path),
        "item_ids": [item["id"] for item in items],
    }
    student_md = student_markdown(manifest, items)
    teacher_md = teacher_markdown(manifest, items)
    (output_dir / "source.md").write_text(
        student_md + "\n\n" + teacher_md,
        encoding="utf-8",
    )
    validation = validate_rendered(manifest, items, student_md)
    if not validation["ok"]:
        (output_dir / "validation.json").write_text(
            json.dumps(validation, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        raise RenderError("组卷校验失败：" + "；".join(validation["errors"]))
    create_docx(output_dir / "学生卷.docx", manifest, items, student=True)
    create_docx(output_dir / "教师卷.docx", manifest, items, student=False)
    create_pdf(output_dir / "学生卷.pdf", manifest, items, student=True)
    create_pdf(output_dir / "教师卷.pdf", manifest, items, student=False)
    formula_cache = output_dir / ".formula-cache"
    if formula_cache.exists():
        shutil.rmtree(formula_cache, ignore_errors=True)
    qa_pages = render_pdf_pngs(output_dir / "学生卷.pdf", output_dir / "qa" / "student")
    teacher_pages = render_pdf_pngs(output_dir / "教师卷.pdf", output_dir / "qa" / "teacher")
    validation["rendered_student_pages"] = len(qa_pages)
    validation["rendered_teacher_pages"] = len(teacher_pages)
    validation["visual_review_required"] = True
    if not qa_pages:
        validation["warnings"].append("未找到 pdftoppm，未能自动生成页面预览图")
    (output_dir / "validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {
        "output": str(output_dir),
        "validation": validation,
        "files": sorted(path.name for path in output_dir.iterdir() if path.is_file()),
    }


def datetime_now() -> str:
    from datetime import datetime

    return datetime.now().astimezone().isoformat(timespec="seconds")


if __name__ == "__main__":
    print("请通过 paperbank.py render 调用此模块。")
