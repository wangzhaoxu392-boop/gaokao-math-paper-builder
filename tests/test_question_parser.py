from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location(
    "question_parser_test_module",
    ROOT / "scripts" / "question_parser.py",
)
parser = importlib.util.module_from_spec(spec)
assert spec.loader
spec.loader.exec_module(parser)


class QuestionParserTests(unittest.TestCase):
    def test_filename_role_precedence(self):
        self.assertEqual(
            parser.correct_role("精品解析：2023年真题（原卷版）.docx")[0],
            "question",
        )
        self.assertEqual(
            parser.correct_role("精品解析：2023年真题（解析版）.docx")[0],
            "solution",
        )
        self.assertEqual(
            parser.correct_role("2026年数学真题（纯答案版）.docx")[0],
            "answer",
        )

    def test_split_sections_options_and_subquestions(self):
        text = """一、选择题：本题共2小题，每小题5分，共10分。
1. 已知集合A，则（ ）
A. 1 B. 2 C. 3 D. 4
2. 已知向量a，则（ ）
A. 1 B. 2 C. 3 D. 4
二、填空题：本题共1小题，每小题5分，共5分。
3. 若x=____，则y=____。
三、解答题：本题共1小题，共10分。
4. 已知数列{a_n}。
（1）求通项；
（2）求前n项和。
"""
        manifest = {
            "source_id": "SRC-TEST",
            "metadata": {
                "grade": "高三",
                "paper_type": "新高考模拟",
                "source_authority": "school_exam",
                "curriculum_version": "new-gaokao-national",
            },
        }
        items = parser.split_question_paper(text, manifest)
        self.assertEqual(len(items), 5)
        self.assertEqual(items[0]["question_type"], "single_choice")
        self.assertEqual(items[0]["options"], ["1", "2", "3", "4"])
        self.assertEqual(items[2]["question_type"], "fill_blank")
        self.assertEqual(items[3]["subquestion_no"], "1")
        self.assertEqual(items[4]["subquestion_no"], "2")
        self.assertEqual(items[3]["primary_band"], "L3")

    def test_compact_answer_table(self):
        answers = parser.split_answer_text("参考答案\n1.C 2.A 3.B 4.D")
        self.assertEqual(answers["1"], "C")
        self.assertEqual(answers["4"], "D")

    def test_compact_answers_keep_multiple_choice_and_later_questions(self):
        text = """参考答案
1.A 2.C 3.D 4.A 5.B 6.B 7.C 8.B
9.BC 10.ACD 11.ABD
12.32
13.ln2
14.0.5
15.(1)B=60°
(2)答案为 2
"""
        parsed = parser.parse_answer_text(text)
        self.assertEqual(parsed["9"]["answer"], "BC")
        self.assertEqual(parsed["11"]["answer"], "ABD")
        self.assertEqual(parsed["12"]["answer"], "32")
        self.assertIn("答案为 2", parsed["15"]["solution"])

    def test_full_solution_separates_answer_and_explanation(self):
        text = """1. 已知集合 A，则（ ）
A. 1 B. 2 C. 3 D. 4
【答案】C
【解析】
【详解】由题意可得 C 正确，故选 C。
"""
        parsed = parser.parse_answer_text(text)
        self.assertEqual(parsed["1"]["answer"], "C")
        self.assertIn("由题意", parsed["1"]["solution"])

    def test_multiple_choice_section_and_decimal_noise(self):
        text = """一、选择题：本题共8小题，每小题5分，共40分。
1. 已知集合A，则（ ）
A. 1 B. 2 C. 3 D. 4
二、选择题：本题共3小题，每小题6分，共18分。本题有多项符合题目要求。
9. 下列结论正确的是（ ）
A. 1 B. 2 C. 3 D. 4
三、填空题：本题共3小题，每小题5分，共15分。
12. 若x=____，则y=____。
0.001
0.010
6.635
"""
        manifest = {"source_id": "SRC-MULTI", "metadata": {"grade": "高三"}}
        items = parser.split_question_paper(text, manifest)
        self.assertTrue(any(item["question_type"] == "multiple_choice" for item in items))
        self.assertFalse(any(int(item["question_no"]) == 0 for item in items))
        self.assertFalse(any(item["question_no"] == 6 and item["stem_markdown"].startswith("6.635") for item in items))

    def test_answer_markers_end_question_and_admin_text_is_ignored(self):
        text = """一、填空题：本题共2小题，每小题5分，共10分。
1. 已知函数f(x)，若f(1)=____，求参数的值。
【答案】3
【解析】
【详解】由已知条件计算可得，所以答案为3。
2. 已知数列a_n，求a_2=____。
【答案】5
参考答案
1. 由已知条件直接计算，答案仍为3。
2. 由递推式计算，答案为5。
绝密★启用前
1.答卷前，考生必须填写姓名。
"""
        manifest = {"source_id": "SRC-BOUNDARY", "metadata": {"grade": "高三"}}
        items = parser.split_question_paper(text, manifest)
        self.assertEqual([item["question_no"] for item in items], [1, 2])
        for item in items:
            self.assertNotIn("【答案】", item["stem_markdown"])
            self.assertNotIn("参考答案", item["stem_markdown"])
            self.assertNotIn("答卷前", item["stem_markdown"])
            self.assertEqual(item["split_rule_version"], 3)

    def test_inline_answer_keeps_clean_stem(self):
        text = """一、选择题：本题共1小题，每小题5分，共5分。
1. 已知集合A={1,2}，则A中元素个数为（ ）
A. 1 B. 2 C. 3 D. 4
【答案】B
【解析】
【详解】集合中有两个元素，故选B。
"""
        manifest = {"source_id": "SRC-INLINE", "metadata": {"grade": "高三"}}
        items = parser.split_question_paper(text, manifest)
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["options"], ["1", "2", "3", "4"])
        self.assertNotIn("【答案】", items[0]["stem_markdown"])
        self.assertEqual(items[0]["split_rule_version"], 3)

    def test_docx_rich_extraction_keeps_embedded_image(self):
        with tempfile.TemporaryDirectory(prefix="gm-docx-parser-") as directory:
            root = Path(directory)
            image_path = root / "formula.png"
            Image.new("RGB", (120, 40), "white").save(image_path)
            document = Document()
            document.add_paragraph("一、选择题：本题共1小题，每小题5分，共5分。")
            paragraph = document.add_paragraph("1. 已知 ")
            paragraph.add_run().add_picture(str(image_path), width=None)
            paragraph.add_run("，则结果为（ ）")
            docx_path = root / "paper.docx"
            document.save(docx_path)
            extracted = parser.extract_docx_rich(docx_path)
            self.assertEqual(extracted["status"], "ok")
            self.assertIn("[[MEDIA:word/media/image1.png]]", extracted["text"])
            self.assertEqual(len(extracted["assets"]), 1)

    def test_paper_without_section_titles(self):
        text = """2026年北京市高考数学
1. 已知集合A，则结果为（ ）
A. 1 B. 2 C. 3 D. 4
2. 已知复数z，则|z|=（ ）
A. 1 B. 2 C. 3 D. 4
3. 若x=____，则y=____。
4. 已知数列，求通项公式。
"""
        manifest = {
            "source_id": "SRC-NO-SECTION",
            "metadata": {"grade": "高三", "paper_type": "新高考真题"},
        }
        items = parser.split_question_paper(text, manifest)
        self.assertEqual(len(items), 4)
        self.assertEqual(items[0]["question_type"], "single_choice")
        self.assertEqual(items[2]["question_type"], "fill_blank")
        self.assertEqual(items[3]["question_type"], "solution")


if __name__ == "__main__":
    unittest.main()
